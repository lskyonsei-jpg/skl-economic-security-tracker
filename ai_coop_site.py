"""패널 2 — 한미 AI 협력 트래커 화면. data/ai_coop_events.json 을 읽어 site/ai_coop.html(Dashboard), site/records.html(사건 목록),
site/review.html(판단 대기), site/ai_coop.json, site/ai_coop_events.xlsx 를 만든다.

사건 파일은 collect_ai_coop.py 가 만든다(seed: 연구자 확정 코딩 CSV, 이후 Task A/B 자동 수집). 파일이 없으면
시드에서 즉석 생성한다. 2026-10-22 Notion 연동은 폐기했다(연구자 결정).

원칙
- 분석 단위는 거래(deal) 하나(스키마 v2, 2026-10-26). count 와 share 를 같이 쓰고 분모는 전체 거래 수다.
- lifecycle(Rumour < Signed < Contracted < Operating < Completed)은 거래가 도달한 단계, canceled 는 0/1. 단계별 날짜 열로 이력을 그린다.
- status(confirmed/auto/rejected) 를 화면에 그대로 보인다. 모델 코딩을 확정 라벨처럼 보이지 않게 한다.
- 금액은 명시된 억 달러만 쓰고 추정치는 합산하지 않는다.
"""
from __future__ import annotations

import html
import json
import re
from collections import Counter
from pathlib import Path

import collect_ai_coop as cc

HERE = Path(__file__).resolve().parent
EVENTS = HERE / "data" / "ai_coop_events.json"
REPO = "lskyonsei-jpg/skl-economic-security-tracker"
STAGES = cc.LIFECYCLES                     # 낮은 단계부터
STAGE_DATE = cc.STAGE_DATE
# 미국 행정부 경계(타임라인 세로선). 취임일 기준.
ADMINISTRATIONS = [("오바마", "2009-01-20"), ("트럼프 1기", "2017-01-20"), ("바이든", "2021-01-20"), ("트럼프 2기", "2025-01-20")]


GOV_RE = re.compile(r"^정부\((.+)\)$")


def disp(name: str) -> str:
    """화면 표기. 자료의 표준 표기 '정부(미국)' 은 그대로 두고(모델 프롬프트·중복 판정이 이 표기를 쓴다) 화면에서만 '미국(정부)' 로 보인다."""
    m = GOV_RE.match(name or "")
    return f"{m.group(1)}(정부)" if m else name


def load_orgs() -> dict:
    """기관 표기 목록(reference/ai_coop_orgs.csv) → {화면 표기: {country, type}}."""
    import csv
    path = HERE / "reference" / "ai_coop_orgs.csv"
    if not path.is_file():
        return {}
    with path.open(encoding="utf-8", newline="") as f:
        return {disp(r["canonical"]): {"country": r.get("country", ""), "type": r.get("type", "")} for r in csv.DictReader(f)}


def load_events() -> tuple[list[dict], dict]:
    """data/ai_coop_events.json -> 화면용 사건 목록. 없으면 시드에서 만든다. rejected 는 뺀다."""
    if not EVENTS.is_file():
        import collect_ai_coop
        if any(collect_ai_coop.SEED_DIR.glob("*.csv")):
            collect_ai_coop.cmd_seed()
    if not EVENTS.is_file():
        return [], {"as_of": ""}
    raw, _ = cc.migrate_doc(json.loads(EVENTS.read_text(encoding="utf-8")))   # Release 자산이 v1 이어도 화면은 v2 로
    events = []
    for e in raw.get("events", []):
        if e.get("status") == "rejected" or not e.get("date"):
            continue
        amt = e.get("amount_usd_bn")
        stage_dates = {col: e.get(col) for col in STAGE_DATE.values()}
        first_stage = next((s for s in STAGES if stage_dates[STAGE_DATE[s]] == e["date"]), e.get("lifecycle"))
        events.append({
            "id": e["id"], "date": e["date"][:10],
            "us": [disp(x) for x in e.get("us_parties", [])], "kr": [disp(x) for x in e.get("kr_parties", [])],
            "actor_type": e.get("track") or "미분류",
            "lifecycle": e.get("lifecycle") or "Rumour", "first_stage": first_stage, "canceled": int(e.get("canceled") or 0),
            **stage_dates, "date_canceled": e.get("date_canceled"),
            "sector": e.get("sectors") or ["미분류"],
            "content": e.get("title", ""), "summary": e.get("summary", ""),
            "amount": (f"{amt:g}억 달러" + (" (추정)" if e.get("amount_is_estimate") else "")) if amt is not None else "",
            "amount_usd_bn": amt, "amount_is_estimate": bool(e.get("amount_is_estimate")),
            "url": (e.get("sources") or [{}])[0].get("url", ""), "sources": [x.get("url", "") for x in (e.get("sources") or []) if x.get("url")],
            "updates": e.get("updates", []), "coded_by": e.get("coded_by", ""),
            "source_grade": (e.get("sources") or [{}])[0].get("grade"),
            "status": e.get("status", "auto"),
            "confidence": e.get("confidence"), "administration": e.get("administration", ""),
            "merged_from": e.get("merged_from") or [], "note": e.get("note", ""),
        })
    events.sort(key=lambda x: (x["date"], x["content"]))
    return events, {"as_of": raw.get("as_of", "")}


def aggregate(events: list[dict]) -> dict:
    n = len(events)
    stages = list(STAGES)
    sectors = sorted({s for e in events for s in e["sector"]})
    years = sorted({e["date"][:4] for e in events})

    by_year_stage = {y: Counter() for y in years}
    for e in events:
        by_year_stage[e["date"][:4]][e["lifecycle"]] += 1
    matrix = {s: Counter() for s in sectors}
    for e in events:
        for s in e["sector"]:
            matrix[s][e["lifecycle"]] += 1
    pairs = Counter()
    for e in events:
        for u in e["us"] or ["(미국 측 미기재)"]:
            for k in e["kr"] or ["(한국 측 미기재)"]:
                pairs[(u, k)] += 1
    actor_type = Counter(e["actor_type"] for e in events)
    # 이행 사다리: 첫 관측 단계 → 현재 단계. 대각선 위가 레코드 안에서 단계가 오른 거래.
    ladder = {a: Counter() for a in stages}
    for e in events:
        ladder[e["first_stage"] or e["lifecycle"]][e["lifecycle"]] += 1
    progressed = sum(1 for e in events if e["first_stage"] and STAGES.index(e["lifecycle"]) > STAGES.index(e["first_stage"]))
    return {
        "n": n, "stages": stages, "sectors": sectors, "years": years,
        "by_year_stage": {y: dict(c) for y, c in by_year_stage.items()},
        "matrix": {s: dict(c) for s, c in matrix.items()},
        "pairs": [{"us": u, "kr": k, "n": c} for (u, k), c in pairs.most_common(30)],
        "actor_type": dict(actor_type),
        "stage_total": dict(Counter(e["lifecycle"] for e in events)),
        "ladder": {a: dict(c) for a, c in ladder.items()}, "progressed": progressed,
        "canceled": sum(e["canceled"] for e in events),
        "status": dict(Counter(e["status"] for e in events)),
        # 구속력 있는 단계(Contracted 이상)의 명시 금액 합계. 추정치 제외
        "amount_contracted_total": round(sum(e["amount_usd_bn"] for e in events
                                             if e["amount_usd_bn"] and not e["amount_is_estimate"] and STAGES.index(e["lifecycle"]) >= STAGES.index("Contracted")), 1),
    }


# ------------------------------------------------------------------ 화면


# 페이지 틀(내비·히어로 띠·Last update·푸터)은 publish_site.page() 가 씌운다. 여기서는 히어로 안쪽과 본문만 만든다.
TITLE = "한미 AI 협력 트래커"


def subnav(current: str) -> str:
    """AI Cooperation 히어로 아래 탭: Dashboard(ai_coop.html) · Records(records.html) · Review(review.html)."""
    tabs = [("ai_coop.html", "Dashboard"), ("records.html", "Records"), ("review.html", "Review")]
    cur = ' aria-current="page"'
    items = "".join(f'<a href="{h}"{cur if h == current else ""}>{t}</a>' for h, t in tabs)
    return f'<nav class="subnav">{items}</nav>'


def hero(current: str, meta: str = "") -> str:
    return (f'<span class="eyebrow">Panel 2 · AI Cooperation</span><div class="hero-row"><h1>{html.escape(TITLE)}</h1>'
            f'{subnav(current)}</div>{meta}')


def render(agg: dict, meta: dict) -> dict:
    """ai_coop.html 의 페이지 사양(publish_site.page 인자 + file)."""
    n = agg["n"]
    asof = meta.get("fetched_at", "")[:10] or "—"
    spec = {"file": "ai_coop.html", "title": f"{TITLE} · SKL", "hero": hero("ai_coop.html"),
            "current": "ai_coop.html", "asof": asof}
    if n == 0:
        return spec | {"body": '<div class="wrap section"><div class="empty">아직 사건 자료가 없습니다.</div></div>'}

    def pct(c: int) -> str:
        return f"{c / n * 100:.1f}%"

    # 매트릭스(분야 x 단계)
    head = "".join(f'<th class="num">{html.escape(f)}</th>' for f in agg["stages"])
    rows = []
    for s in agg["sectors"]:
        cells = "".join(
            f'<td class="cell">{agg["matrix"][s].get(f, 0)}<small>{pct(agg["matrix"][s].get(f, 0))}</small></td>'
            for f in agg["stages"])
        total = sum(agg["matrix"][s].values())
        rows.append(f'<tr><td>{html.escape(s)}</td>{cells}<td class="cell"><b>{total}</b><small>{pct(total)}</small></td></tr>')
    matrix_html = (f'<table class="matrix"><thead><tr><th>분야</th>{head}<th class="num">합계</th></tr></thead>'
                   f'<tbody>{"".join(rows)}</tbody></table>'
                   f'<p class="note">건수 · 전체 {n}건 대비 비중. 여러 분야에 걸친 거래는 분야마다 집계. 단계는 현재 도달 단계(lifecycle).</p>')

    # 단계별 막대
    bars = "".join(
        f'<div class="row"><span>{html.escape(f)}</span><div><div class="bar" style="width:{agg["stage_total"].get(f, 0) / n * 100:.1f}%"></div></div>'
        f'<span class="n">{agg["stage_total"].get(f, 0)} · {pct(agg["stage_total"].get(f, 0))}</span></div>' for f in agg["stages"])
    bars += (f'<p class="note">현재 도달 단계 기준. 취소 {agg["canceled"]}건은 단계를 유지한 채 canceled=1. '
             f'Contracted 이상의 명시 금액 합계 {agg["amount_contracted_total"]:g}억 달러(추정치 제외).</p>')

    # 이행 사다리: 첫 관측 단계(행) → 현재 단계(열)
    lhead = "".join(f'<th class="num">{html.escape(f)}</th>' for f in agg["stages"])
    lrows = []
    for a in agg["stages"]:
        row_total = sum(agg["ladder"][a].values())
        cells = "".join(f'<td class="cell{" up" if agg["stages"].index(b) > agg["stages"].index(a) and agg["ladder"][a].get(b, 0) else ""}">{agg["ladder"][a].get(b, 0) or ""}</td>' for b in agg["stages"])
        lrows.append(f'<tr><td>{html.escape(a)}</td>{cells}<td class="cell"><b>{row_total}</b></td></tr>')
    ladder_html = (f'<table class="matrix"><thead><tr><th>첫 관측 ↓ · 현재 →</th>{lhead}<th class="num">합계</th></tr></thead><tbody>{"".join(lrows)}</tbody></table>'
                   f'<p class="note">대각선 위(색칠)가 레코드 안에서 단계가 오른 거래 {agg["progressed"]}건. 거래는 처음 보도된 단계로 들어오므로 '
                   f'Contracted 에서 시작한 거래는 그 앞 단계가 관측되지 않은 것이지 건너뛴 것이 아니다. 최근 거래는 아직 오를 시간이 없었다(우측 절단).</p>')

    pairs = "".join(
        f'<li><span>{html.escape(p["us"])}</span><span class="arrow">↔</span><span>{html.escape(p["kr"])}</span><span class="n">{p["n"]}</span></li>'
        for p in agg["pairs"][:15])

    body = f'''
<div class="wrap">
  <section class="section">
    <span class="eyebrow">Timeline</span>
    <div class="panel"><div class="panel-head"><span>거래 타임라인 — 단계별 날짜, 같은 거래는 선으로 연결</span><span id="tl-legend" class="legend"></span><a href="codebook.html#panel2">코드북</a></div>
      <div class="panel-body timeline-wrap"><svg id="timeline" viewBox="0 0 1000 440" aria-label="한미 AI 협력 거래 타임라인"></svg><div id="tl-tip" class="tooltip"></div></div></div>
  </section>
  <section class="section grid grid-2">
    <div><span class="eyebrow">Lifecycle</span><div class="panel"><div class="panel-head">단계별 건수 · 비중 (분모 {n}건)</div><div class="panel-body bars">{bars}</div></div></div>
    <div><span class="eyebrow">Ladder</span><div class="panel"><div class="panel-head">이행 사다리 — 첫 관측 단계 → 현재 단계</div><div class="panel-body" style="overflow-x:auto">{ladder_html}</div></div></div>
  </section>
  <section class="section grid grid-2">
    <div><span class="eyebrow">Sector × Lifecycle</span><div class="panel"><div class="panel-head">분야 × 단계 매트릭스</div><div class="panel-body" style="overflow-x:auto">{matrix_html}</div></div></div>
    <div><span class="eyebrow">Actors</span><div class="panel"><div class="panel-head">미국 주체 ↔ 한국 주체 (상위 15 조합)</div><div class="panel-body"><ul class="pairs" style="list-style:none;margin:0;padding:0">{pairs}</ul></div></div></div>
  </section>
</div>
<script>
const STAGES={json.dumps(agg["stages"], ensure_ascii=False)};
const STAGE_DATE={json.dumps(STAGE_DATE, ensure_ascii=False)};
const ADMIN={json.dumps(ADMINISTRATIONS, ensure_ascii=False)};
const css=v=>getComputedStyle(document.documentElement).getPropertyValue(v).trim();
const STAGE_COLOR={{}};STAGES.forEach((f,i)=>STAGE_COLOR[f]=css('--cat-'+(i+1)));
fetch('ai_coop.json',{{cache:'no-store'}}).then(r=>r.json()).then(D=>{{
  const E=D.events; const tip=document.getElementById('tl-tip');
  document.getElementById('tl-legend').innerHTML=STAGES.map(f=>`<span><i style="background:${{STAGE_COLOR[f]}}"></i>${{f}}</span>`).join('')+`<span><i style="background:transparent;border:2px solid ${{css('--ink')}}"></i>canceled</span>`;
  const svg=document.getElementById('timeline'); const W=1000,H=440,L=150,R=24,T=44,B=50;
  const allDates=E.flatMap(e=>STAGES.map(s=>e[STAGE_DATE[s]]).filter(Boolean)).sort();
  const y0=+allDates[0].slice(0,4), y1=+allDates[allDates.length-1].slice(0,4)+1;
  const x=d=>L+(new Date(d+'T00:00:00Z')-Date.UTC(y0,0,1))/(Date.UTC(y1,0,1)-Date.UTC(y0,0,1))*(W-L-R);
  let s='';
  for(let y=y0;y<=y1;y++){{const xx=x(y+'-01-01');s+=`<line x1="${{xx}}" x2="${{xx}}" y1="${{T}}" y2="${{H-B}}" stroke="${{css('--line')}}"/>`; if((y-y0)%2===0||y1-y0<8) s+=`<text class="tl-axis" x="${{xx+4}}" y="${{H-B+26}}">${{y}}</text>`;}}
  ADMIN.forEach(([name,d])=>{{const yy=d.slice(0,4); if(yy<y0||yy>y1)return; const xx=x(d); s+=`<line x1="${{xx}}" x2="${{xx}}" y1="${{T-14}}" y2="${{H-B}}" stroke="${{css('--ink')}}" stroke-dasharray="3 3"/><text class="tl-admin" x="${{xx+6}}" y="${{T-2}}">${{name}}</text>`;}});
  const lanes=STAGES.length, laneH=(H-T-B-20)/Math.max(lanes,1);
  STAGES.forEach((f,i)=>{{const yy=T+20+laneH*i+laneH/2; s+=`<text class="tl-axis" x="${{L-10}}" y="${{yy+8}}" text-anchor="end">${{f}}</text>`;}});
  // 거래마다 단계 날짜가 있는 칸에 점을 찍고 같은 거래의 점은 선으로 잇는다. 같은 줄에서 가로로 16 단위 안에 있는 점은
  // 위아래 칸(±16)으로 비켜 놓는다 — 점마다 따로 집을 수 있게 (연구자 지시 2026-10-23)
  const placed=STAGES.map(()=>[]); const STEP=16, MINDX=16; const pts=[];
  E.forEach((e,i)=>STAGES.forEach((st,li)=>{{const d=e[STAGE_DATE[st]]; if(d) pts.push({{e,i,st,li,cx:x(d),d}});}}));
  pts.sort((a,b)=>a.cx-b.cx).forEach(p=>{{
    const order=[0,1,-1,2,-2,3,-3]; let k=order.find(k=>!placed[p.li].some(q=>q.k===k&&Math.abs(q.cx-p.cx)<MINDX)); if(k===undefined)k=0;
    placed[p.li].push({{k,cx:p.cx}}); p.cy=T+20+laneH*p.li+laneH/2+k*STEP;}});
  const byDeal={{}}; pts.forEach(p=>(byDeal[p.i]=byDeal[p.i]||[]).push(p));
  Object.values(byDeal).forEach(ps=>{{ps.sort((a,b)=>a.li-b.li); for(let j=1;j<ps.length;j++) s+=`<line x1="${{ps[j-1].cx}}" y1="${{ps[j-1].cy}}" x2="${{ps[j].cx}}" y2="${{ps[j].cy}}" stroke="${{css('--ink-2')}}" stroke-width="1.5" stroke-opacity=".6"/>`;}});
  pts.forEach(p=>{{const cur=p.st===p.e.lifecycle; const fill=STAGE_COLOR[p.st]||css('--muted');
    s+=`<circle data-i="${{p.i}}" data-st="${{p.st}}" data-d="${{p.d}}" cx="${{p.cx}}" cy="${{p.cy}}" r="${{cur?7:5}}" fill="${{p.e.canceled?'transparent':fill}}" stroke="${{p.e.canceled?css('--ink'):fill}}" stroke-width="2" fill-opacity="${{cur?'.95':'.55'}}" style="cursor:pointer"/>`;}});
  svg.innerHTML=s;
  svg.addEventListener('mousemove',ev=>{{const c=ev.target.closest('circle'); if(!c){{tip.style.display='none';return;}} const e=E[+c.dataset.i]; tip.innerHTML=`<b>${{c.dataset.d}}</b> ${{c.dataset.st}}${{e.canceled?' · canceled':''}}<br>${{e.content}}<br>${{e.us.join(', ')}} ↔ ${{e.kr.join(', ')}}`; const r=svg.parentElement.getBoundingClientRect(); tip.style.left=(ev.clientX-r.left+12)+'px'; tip.style.top=(ev.clientY-r.top+12)+'px'; tip.style.display='block';}});
  svg.addEventListener('mouseleave',()=>tip.style.display='none');
  svg.addEventListener('click',ev=>{{const c=ev.target.closest('circle'); if(c) location.href='records.html#'+encodeURIComponent(E[+c.dataset.i].id);}});
}});
</script>'''
    return spec | {"body": body}

# ------------------------------------------------------------------ 사건 목록(records.html) · Excel

# Excel 열 순서. 레코드에 이 밖의 키가 생기면 뒤에 이름순으로 붙는다(변수를 빼지 않는다).
XLSX_ORDER = ["id", "date", "title", "lifecycle", "canceled", "date_rumour", "date_signed", "date_contracted", "date_operating", "date_completed", "date_canceled",
              "us_parties", "kr_parties", "track", "track_original", "sectors", "amount_usd_bn", "amount_is_estimate", "amount_krw_bn",
              "summary", "note", "status", "administration", "source_grade", "confidence", "coded_by", "coded_at", "model",
              "merged_from", "sources", "updates", "schema_version"]
XLSX_FILE = "ai_coop_events.xlsx"
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _xlsx_cell(v):
    """레코드 값 → 셀 값. 목록은 '; ' 로 잇고, 출처는 url(grade), updates 등 구조는 JSON 문자열로 둔다."""
    if v is None or isinstance(v, (int, float, bool)):
        return v
    if isinstance(v, list):
        if all(isinstance(x, str) for x in v):
            return "; ".join(v)
        if all(isinstance(x, dict) and "url" in x for x in v):
            return "; ".join(f'{x.get("url", "")}' + (f' ({x["grade"]})' if x.get("grade") is not None else "") for x in v)
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return _ILLEGAL.sub("", str(v))


def write_xlsx(out: Path, as_of: str) -> int:
    """data/ai_coop_events.json 의 사건 레코드를 변수 그대로 site/ai_coop_events.xlsx 에 쓴다(rejected 제외). 반환값은 행 수."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    raw, _ = cc.migrate_doc(json.loads(EVENTS.read_text(encoding="utf-8"))) if EVENTS.is_file() else ({}, False)
    rows = [e for e in raw.get("events", []) if e.get("status") != "rejected" and e.get("date")]
    rows.sort(key=lambda e: (e["date"], e.get("title", "")), reverse=True)
    keys = {k for e in rows for k in e}
    cols = [c for c in XLSX_ORDER if c in keys] + sorted(keys - set(XLSX_ORDER))
    wb = Workbook()
    ws = wb.active
    ws.title = "events"
    ws.append(cols)
    for r, e in enumerate(rows, start=2):
        for j, c in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=j, value=_xlsx_cell(e.get(c)))
            if isinstance(cell.value, str):
                cell.data_type = "s"   # '=' 로 시작하는 문자열을 수식으로 읽지 않게
    ws.freeze_panes = "A2"
    for j, c in enumerate(cols, start=1):
        width = max([len(c)] + [min(len(str(ws.cell(row=r, column=j).value or "")), 60) for r in range(2, len(rows) + 2)])
        ws.column_dimensions[get_column_letter(j)].width = max(8, min(width + 2, 62))
    meta = wb.create_sheet("meta")
    for row in (("file", XLSX_FILE), ("source", "data/ai_coop_events.json"), ("as_of", as_of), ("n_events", len(rows)),
                ("excluded", "status == rejected"), ("lists", "목록 값은 '; ' 로 이어 붙임. sources 는 url (grade). updates 는 JSON 문자열."),
                ("codebook", "codebook.html#panel2")):
        meta.append(row)
    meta.column_dimensions["A"].width = 12
    meta.column_dimensions["B"].width = 80
    wb.save(out / XLSX_FILE)
    return len(rows)


# 표 열. (변수명, 셀 class). 값도 자료 그대로 보인다(한국어 라벨로 바꾸지 않는다 — 연구자 지시 2026-10-26).
RECORD_COLS = [("date", "d"), ("title", ""), ("lifecycle", ""), ("canceled", "num"), ("sectors", ""), ("track", ""), ("us_parties", ""), ("kr_parties", ""),
               ("amount_usd_bn", "num"), ("status", "")]


def render_records(agg: dict, meta: dict) -> dict:
    """records.html 의 페이지 사양. 사건 목록 전체 표 + 검색 + 기관 현황 + Excel 다운로드."""
    n = agg["n"]
    asof = meta.get("fetched_at", "")[:10] or "—"
    spec = {"file": "records.html", "title": f"거래 목록 · {TITLE} · SKL", "hero": hero("records.html"),
            "current": "records.html", "asof": asof}
    if n == 0:
        return spec | {"body": '<div class="wrap section"><div class="empty">아직 사건 자료가 없습니다.</div></div>'}
    head = "".join(f'<th{(" class=" + chr(34) + c + chr(34)) if c else ""}>{html.escape(k)}</th>' for k, c in RECORD_COLS)
    body = f'''
<div class="wrap">
  <section class="section">
    <span class="eyebrow">Records</span>
    <div class="panel">
      <div class="panel-head records-tools"><span>거래 목록 {n}건 (최신순)</span>
        <span class="tools"><input id="q" type="search" placeholder="검색"><a class="btn small" href="{XLSX_FILE}" download>Excel 다운로드</a><a href="codebook.html#panel2">코드북</a></span></div>
      <div class="tablebox" style="border:0;margin:0"><div id="org-card"></div>
        <table id="records"><thead><tr>{head}</tr></thead><tbody></tbody></table></div>
    </div>
    <p class="note">레코드 하나 = 거래 하나. 열 이름은 자료 변수명 그대로. title 을 누르면 나머지 변수(id · 단계별 날짜 date_rumour…date_completed · date_canceled · administration · summary · note · merged_from · updates · sources · coded_by · confidence · source_grade)가 열린다. 기관명을 누르면 그 기관의 거래만 남는다.
    Excel 은 <code>data/ai_coop_events.json</code> 의 레코드를 변수 전부 그대로 담는다(rejected 제외 · 목록 값은 "; " 로 이어 붙임 · as_of {html.escape(asof)}).</p>
  </section>
</div>
<script>
const STAGES={json.dumps(agg["stages"], ensure_ascii=False)};
const STAGE_DATE={json.dumps(STAGE_DATE, ensure_ascii=False)};
fetch('ai_coop.json',{{cache:'no-store'}}).then(r=>r.json()).then(D=>{{
  const E=D.events; const tb=document.querySelector('#records tbody'); const esc=t=>String(t??'').replace(/[&<>"]/g,c=>({{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}})[c]);
  const ORGS=D.orgs||{{}}; let orgSel=null;
  const orgLink=n=>`<a href="#" class="org" data-org="${{esc(n)}}">${{esc(n)}}</a>`;
  const amt=e=>e.amount_usd_bn==null?'':esc(e.amount_usd_bn)+(e.amount&&e.amount.includes('추정')?' <span class="small">(estimate)</span>':'');
  const detail=e=>{{const ups=(e.updates||[]).map(u=>`<li>${{esc(u.date||'')}} ${{esc(u.field||u.kind||'')}}${{u.old!==undefined?` ${{esc(u.old)}} → ${{esc(u.new)}}`:''}}${{u.evidence?` <span class="small">${{esc(u.evidence)}}</span>`:''}}</li>`).join('');
    const srcs=(e.sources||[]).map((u,i)=>`<a href="${{esc(u)}}" target="_blank" rel="noopener">${{esc(u)}}</a>`).join('<br>');
    const row=(k,v)=>v?`<span>${{k}}</span><span>${{v}}</span>`:'';
    const stages=STAGES.map(st=>e[STAGE_DATE[st]]?`${{st}} ${{esc(e[STAGE_DATE[st]])}}`:null).filter(Boolean).join(' → ');
    return `<div class="ev-detail"><div class="ev-grid">`+row('id',esc(e.id))+row('stages',stages)+row('date_canceled',esc(e.date_canceled))+row('administration',esc(e.administration))
      +row('summary',esc(e.summary))+row('note',esc(e.note).replace(/\\n/g,'<br>'))
      +row('merged_from',(e.merged_from||[]).map(esc).join(', '))
      +row('updates',ups?`<ul class="ev-updates">${{ups}}</ul>`:'')+row('sources',srcs)
      +row('coded_by',esc(e.coded_by))+row('confidence',e.confidence==null?'':esc(e.confidence))+row('source_grade',e.source_grade==null?'':esc(e.source_grade))+`</div></div>`;}};
  const draw=q=>{{const qq=(q||'').toLowerCase(); const rows=[...E].reverse().filter(e=>(!orgSel||e.us.includes(orgSel)||e.kr.includes(orgSel))&&(!qq||JSON.stringify(e).toLowerCase().includes(qq)));
    tb.innerHTML=rows.map((e,i)=>`<tr class="ev-row" data-i="${{i}}" data-id="${{esc(e.id)}}"><td class="d">${{e.date}}</td><td><a href="#" class="ev-title">${{esc(e.content)}}</a></td><td>${{esc(e.lifecycle)}}</td><td class="num">${{e.canceled?1:0}}</td><td>${{esc(e.sector.join(', '))}}</td><td>${{esc(e.actor_type)}}</td><td>${{e.us.map(orgLink).join(', ')}}</td><td>${{e.kr.map(orgLink).join(', ')}}</td><td class="num">${{amt(e)}}</td><td>${{esc(e.status)}}</td></tr><tr class="ev-more" hidden><td colspan="${{{len(RECORD_COLS)}}}">${{detail(e)}}</td></tr>`).join('');
    const card=document.getElementById('org-card');
    if(orgSel){{const o=ORGS[orgSel]||{{}}; const st={{}}; rows.forEach(e=>st[e.lifecycle]=(st[e.lifecycle]||0)+1); const partners={{}}; rows.forEach(e=>(e.us.includes(orgSel)?e.kr:e.us).forEach(p=>partners[p]=(partners[p]||0)+1));
      card.innerHTML=`<div class="org-box"><b>${{esc(orgSel)}}</b> <span class="small">${{o.type?esc(o.type)+' · ':''}}${{o.country?esc(o.country)+' · ':''}}거래 ${{rows.length}}건 · ${{STAGES.filter(f=>st[f]).map(f=>f+' '+st[f]).join(' · ')}}</span><br><span class="small">상대: ${{Object.entries(partners).sort((a,b)=>b[1]-a[1]).map(([p,n])=>orgLink(p)+' '+n).join(', ')||'-'}}</span> <button class="btn small" id="org-clear" type="button">전체 보기</button></div>`;
      document.getElementById('org-clear').onclick=()=>{{orgSel=null;draw(q);}};}} else card.innerHTML='';}};
  tb.addEventListener('click',ev=>{{const t=ev.target.closest('a.ev-title'); if(t){{ev.preventDefault(); const more=t.closest('tr').nextElementSibling; more.hidden=!more.hidden; return;}}
    const o=ev.target.closest('a.org'); if(o){{ev.preventDefault(); orgSel=o.dataset.org; draw(document.getElementById('q').value); document.getElementById('org-card').scrollIntoView({{block:'nearest'}});}}}});
  document.getElementById('org-card').addEventListener('click',ev=>{{const o=ev.target.closest('a.org'); if(o){{ev.preventDefault(); orgSel=o.dataset.org; draw(document.getElementById('q').value);}}}});
  draw(''); const q=document.getElementById('q');
  const goHash=()=>{{const h=decodeURIComponent(location.hash.slice(1)); if(!h)return; const tr=tb.querySelector(`tr[data-id="${{h.replace(/"/g,'')}}"]`); if(!tr)return;
    tr.nextElementSibling.hidden=false; tb.querySelectorAll('tr.hl').forEach(r=>r.classList.remove('hl')); tr.classList.add('hl'); tr.scrollIntoView({{block:'center'}});}};
  goHash(); window.addEventListener('hashchange',goHash); q.addEventListener('input',ev=>draw(ev.target.value));
}});
</script>'''
    return spec | {"body": body}


# ------------------------------------------------------------------ 판단 대기(review.html)

STATE_FILE = HERE / "data" / "ai_coop_state.json"


def pending_proposals() -> list[dict]:
    if not STATE_FILE.is_file():
        return []
    st, _ = cc.migrate_state(json.loads(STATE_FILE.read_text(encoding="utf-8")))   # Release 자산이 v1 이어도 카드는 v2 로
    return [p for p in st.get("proposals", []) if p.get("status") == "proposed"]


def _esc(v) -> str:
    return html.escape("" if v is None else str(v))


def render_review(props: list[dict], events_by_id: dict) -> dict:
    """review.html 의 페이지 사양(publish_site.page 인자 + file)."""
    by_task = {k: [p for p in props if p["task"] == k] for k in ("A", "B", "C")}
    hero_meta = (f'<div class="meta"><div>판단 대기<b>{len(props)}건</b></div><div>A 새 사건<b>{len(by_task["A"])}</b></div>'
            f'<div>B 갱신<b>{len(by_task["B"])}</b></div><div>C 재검토<b>{len(by_task["C"])}</b></div></div>')
    def ev_line(ev: dict) -> str:
        return (f'{_esc(ev.get("date"))} · {_esc(ev.get("lifecycle"))}{" · canceled" if ev.get("canceled") in (1, True) else ""} · {_esc(ev.get("track"))} · '
                f'{_esc(", ".join(ev.get("sectors", [])))} · 금액 {_esc(ev.get("amount_usd_bn"))} · conf {_esc(ev.get("confidence"))}')
    FIELD_KO = {"amount_usd_bn": "금액(억 달러)", "amount_krw_bn": "금액(억 원)", "lifecycle": "단계(lifecycle)", "canceled": "취소(canceled)", "track": "track",
                "sectors": "분야", "date": "날짜", "title": "제목", "us_parties": "미국 측", "kr_parties": "한국 측", "summary": "요약",
                "date_rumour": "date_rumour", "date_signed": "date_signed", "date_contracted": "date_contracted", "date_operating": "date_operating",
                "date_completed": "date_completed", "date_canceled": "date_canceled"}
    def _lab(field) -> str:
        return FIELD_KO.get(field, _esc(field))
    def _old(v) -> str:   # 기존 판단 값. 비어 있으면 '없음'
        return "없음" if v in (None, "", [], "null") else _esc(", ".join(v) if isinstance(v, list) else v)
    def _new(v, field=None) -> str:   # 모델 판단 값. 그 항목을 비우자는 제안은 '값 삭제 권장'
        return "값 삭제 권장" if v in (None, "", [], "null") else _esc(", ".join(v) if isinstance(v, list) else v)
    def ev_link(tgt: dict) -> str:   # 기존 사건 → 사건 목록(records.html)의 그 행으로 (a 태그라 Ctrl-클릭·가운데 클릭으로 새 탭)
        if not tgt:
            return "기존: -"
        return f'기존: <a class="ev-link" href="records.html#{_esc(tgt.get("id"))}" title="사건 목록에서 이 사건 보기">{ev_line(tgt)}</a>'
    def card(p: dict, title: str, meta: str, body: str) -> str:
        """한 줄 요약(제목·요지) + 결정 라디오를 항상 보이고, 상세는 '열기' 뒤에 둔다."""
        e = p.get("evidence", {})
        link = f'<a href="{_esc(e.get("url"))}" target="_blank" rel="noopener">원문</a>' if e.get("url") else ""
        pid = _esc(p["id"])
        return (f'<div class="panel rcard" data-id="{pid}"><div class="rhead"><span class="small">{pid}{" · " + link if link else ""}</span>'
                f'<b class="rtitle">{title}</b><span class="small rmeta">{meta}</span></div>'
                f'<div class="decide" data-id="{pid}">'
                f'<label><input type="radio" name="d-{pid}" value="approve"> 승인</label>'
                f'<label><input type="radio" name="d-{pid}" value="reject"> 기각</label>'
                f'<label><input type="radio" name="d-{pid}" value="hold"> 보류</label>'
                f'<label><input type="radio" name="d-{pid}" value="other"> 기타</label>'
                f'<input type="text" class="other-note" data-note="{pid}" placeholder="의견·수정 지시 (예: 금액은 20으로, 형태는 집행, 주체에 기아 추가)" hidden>'
                + (f'<span class="small hold-note">어제 의견: {_esc(p.get("hold_note"))}</span>' if p.get("hold_note") else "")
                + (f'<span class="small hold-note">반영 실패: {_esc(p.get("apply_error"))}</span>' if p.get("apply_error") else "") +
                f'<details class="rmore"><summary><span class="btn small"></span></summary>{body}'
                f'<p class="note">원문 발췌({_esc(e.get("via", ""))} · {_esc(e.get("chars", ""))}자): {_esc(str(e.get("excerpt", ""))[:300])}…</p>'
                f'</details></div></div>')
    sections = []
    titles = {"A": "새 사건 후보", "B": "기존 사건 갱신 후보", "C": "기존 판단 재검토"}
    for k in ("A", "B", "C"):
        items = by_task[k]
        cards = []
        for p in items:
            if k == "A":
                ev = p["payload"]
                extra = (f'<p class="small">기존 사건과 중복 가능: {_esc(p.get("possible_duplicate_of"))}</p>' if p.get("possible_duplicate_of") else "") + \
                        (f'<p class="small">기관 목록(reference/ai_coop_orgs.csv)에 없는 표기: {_esc(", ".join(p.get("unknown_orgs", [])))}</p>' if p.get("unknown_orgs") else "")
                more = p.get("extra_sources") or []
                if more:
                    extra += '<p class="small">같은 거래를 다룬 다른 기사 ' + str(len(more)) + '건: ' + " · ".join(
                        f'<a href="{_esc(m.get("url"))}" target="_blank" rel="noopener">{_esc((m.get("title") or "")[:30])}</a>' + (f' ({_esc(m.get("lifecycle"))}, {_esc(m.get("confidence"))})' if m.get("lifecycle") else '') for m in more) + '</p>'
                if p.get("lifecycle_variants"):
                    extra += f'<p class="small">다른 기사의 단계 판정: {_esc(", ".join(sorted(set(str(v) for v in p["lifecycle_variants"]))))}</p>'
                title, meta = _esc(ev.get("title")), ev_line(ev)
                body = (f'<p>미국 측 {_esc(", ".join(disp(x) for x in ev.get("us_parties", [])))} / 한국 측 {_esc(", ".join(disp(x) for x in ev.get("kr_parties", [])))}</p>'
                        f'<p>{_esc(ev.get("summary"))}</p><p class="note">모델 메모: {_esc(ev.get("note"))}</p>{extra}')
            else:  # B 갱신 · C 재검토 — 설명 세 부분(왜 다시 보나 / 기존 ↔ 검색 결과 / 권장 수정안). 2026-10-23
                pl = p["payload"]; tgt = events_by_id.get(p.get("target_id"), {}); ex = p.get("explain") or {}
                if k == "B":
                    title = f'갱신: {_esc(p.get("target_title"))}'
                    what = ("새 거래 추가" if pl.get("new_event") else
                            (f'단계 → {_esc(pl.get("lifecycle"))}' + (f' ({_esc(pl.get("stage_date"))})' if pl.get("stage_date") else "") if pl.get("lifecycle") else
                             ("취소(canceled=1)" if pl.get("canceled") else "필드 갱신")))
                    meta = f'{what} · conf {_esc(p.get("confidence"))}'
                else:
                    title = f'재검토: {_esc(p.get("target_title"))}'
                    meta = ""   # 표와 중복이라 비움 (연구자 지시 2026-10-23)
                rec = []
                if pl.get("new_event"):
                    rec.append(f'<p>새 거래 추가(같은 당사자의 별개 거래): <b>{_esc(pl["new_event"].get("title"))}</b> — {ev_line(pl["new_event"])}</p>')
                for u in pl.get("updates", []):
                    rec.append(f'<p><b>{_lab(u.get("field"))}</b> 기존 판단 {_old(u.get("old"))} → 모델 판단 {_new(u.get("new"), u.get("field"))}<br><span class="small">{_esc(u.get("evidence"))}</span></p>')
                if pl.get("lifecycle"):
                    rec.append(f'<p>단계: {_esc(tgt.get("lifecycle") or "Rumour")} → <b>{_esc(pl["lifecycle"])}</b>' + (f' · 도달일 {_esc(pl.get("stage_date"))}' if pl.get("stage_date") else "") + '</p>')
                if pl.get("canceled"):
                    rec.append(f'<p>취소: canceled 0 → <b>1</b>' + (f' · 보도일 {_esc(pl.get("stage_date"))}' if pl.get("stage_date") else "") + f' (단계 {_esc(tgt.get("lifecycle"))} 유지)</p>')
                if pl.get("note_v1_lifecycle"):
                    rec.append(f'<p class="small">변환 전(v1) 제안 값: 진행 상태 {_esc(pl["note_v1_lifecycle"])}</p>')
                if pl.get("diffs"):
                    rows = "".join(f'<tr><td>{_lab(d.get("field"))}</td><td>{_old(d.get("old"))}</td><td>{_new(d.get("new"), d.get("field"))}</td><td class="small">{_esc(d.get("reason"))}</td></tr>' for d in pl["diffs"])
                    rec.append(f'<table><thead><tr><th>항목</th><th>기존 판단</th><th>모델 판단</th><th>기사 근거</th></tr></thead><tbody>{rows}</tbody></table>')
                if ex:
                    body = (f'<div class="explain"><p><b>1. 왜 다시 보나</b><br>{_esc(ex.get("reason")) or "—"}</p>'
                            f'<p><b>2. 기존 코딩 ↔ 검색 결과</b><br>{ev_link(tgt)}<br>검색: {_esc(ex.get("found")) or "—"}</p>'
                            f'<p><b>3. 권장 수정안</b><br>{_esc(ex.get("recommendation")) or "—"}</p>{"".join(rec)}</div>')
                else:  # 2026-10-23 이전 방식으로 만들어진 제안(설명 세 부분 없음)
                    if pl.get("diffs"):
                        rows = "".join(f'<tr><td>{_lab(d.get("field"))}</td><td>{_old(d.get("old"))}</td><td>{_new(d.get("new"), d.get("field"))}</td></tr>' for d in pl["diffs"])
                        rec = [f'<table><thead><tr><th>항목</th><th>기존 판단</th><th>모델 판단</th></tr></thead><tbody>{rows}</tbody></table>']
                    body = f'<p>{ev_link(tgt)}</p>{"".join(rec)}<p class="note">모델 메모: {_esc(p.get("note"))}</p>'
            cards.append(card(p, title, meta, body))
        toggle = '<button class="btn small open-all" type="button">모두 열기</button>' if cards else ''
        sections.append(f'<section class="section rsec"><div class="wrap"><div class="rsec-head"><span class="eyebrow">{k}. {titles[k]} ({len(items)}건)</span>{toggle}</div>'
                        + ("".join(cards) if cards else '<div class="empty">없음</div>') + '</div></section>')
    howto = ('<div class="wrap howto"><details><summary>제출 방법</summary><p class="note">'
             '카드마다 승인·기각·보류를 고르고, 고칠 내용이 있으면 의견란에 적습니다. "GitHub으로 제출"을 누르면 내용이 채워진 GitHub Issue 화면이 열리고, 거기서 등록 버튼만 누르면 됩니다(GitHub 로그인 필요). '
             '다음 날 06:00 자동 실행이 반영합니다. 승인은 그대로, 의견에 수정 지시가 있으면 고쳐서 반영하고, 기각은 기록만 남깁니다. 선택하지 않은 항목은 이월 및 축적됩니다. '
             '선택은 이 브라우저에 임시 저장됩니다.</p></details></div>')
    submit_bar = ('<div class="submit-bar"><span id="decided-count">선택 0건</span><span class="hint">제출하면 선택이 초기화됩니다</span>'
                  '<span><button class="btn" id="clear-decisions" type="button">초기화</button> '
                  '<button class="btn solid" id="submit-decisions" type="button">GitHub으로 제출</button></span></div>')
    script = """
<script>
(function(){
  const KEY='aicoop-decisions-v1'; const REPO=%s;
  const load=()=>{try{return JSON.parse(localStorage.getItem(KEY)||'{}')}catch(e){return {}}};
  const save=d=>{try{localStorage.setItem(KEY,JSON.stringify(d))}catch(e){}};
  let D=load();
  const cards=[...document.querySelectorAll('.decide')];
  const count=()=>{const n=Object.values(D).filter(x=>x.decision).length;document.getElementById('decided-count').textContent='선택 '+n+'건';};
  cards.forEach(c=>{const id=c.dataset.id; const d=D[id]||{};
    const ta=c.querySelector('input.other-note'); const showNote=on=>{ta.hidden=!on; if(on)ta.focus();};
    // 의견란: 기타(보류+의견)와 승인(수정 지시가 있으면 고쳐서 반영 — 코드북 0단계)에서 연다
    c.querySelectorAll('input[type=radio]').forEach(r=>{if(d.decision===r.value)r.checked=true; r.addEventListener('change',()=>{D[id]=Object.assign(D[id]||{},{decision:r.value});save(D);count();showNote(r.value==='other'||r.value==='approve'||!!(D[id].note&&D[id].note.trim()));});});
    if(d.note)ta.value=d.note; if(d.decision==='other'||d.decision==='approve'||d.note)ta.hidden=false; ta.addEventListener('input',()=>{D[id]=Object.assign(D[id]||{},{note:ta.value});save(D);});
  });
  count();
  document.querySelectorAll('.open-all').forEach(b=>b.addEventListener('click',()=>{const ds=[...b.closest('.rsec').querySelectorAll('details.rmore')];const open=ds.some(d=>!d.open);ds.forEach(d=>d.open=open);b.textContent=open?'모두 닫기':'모두 열기';}));
  const word=x=>x==='other'?'hold':(x||'hold');   // 기타 = 보류 + 의견 (자동 반영 없음, 다음 날 카드에 의견이 남는다)
  const lines=()=>Object.entries(D).filter(([id,d])=>d.decision||(d.note&&d.note.trim())).map(([id,d])=>'- '+id+': '+word(d.decision)+(d.note&&d.note.trim()?' | 의견: '+d.note.trim().replace(/\\s*\\n+\\s*/g,' / '):''));
  document.getElementById('submit-decisions').addEventListener('click',()=>{const L=lines(); if(!L.length){alert('선택된 결정이 없습니다.');return;}
    const today=new Date().toISOString().slice(0,10); const body='## 결정 '+today+'\\n\\n'+L.join('\\n')+'\\n\\n(판단 대기 페이지에서 제출. 다음 자동 실행 때 반영됩니다.)';
    const url='https://github.com/'+REPO+'/issues/new?labels=ai-coop-decision&title='+encodeURIComponent('[결정] '+today)+'&body='+encodeURIComponent(body);
    window.open(url,'_blank'); D={}; save(D); cards.forEach(c=>{c.querySelectorAll('input[type=radio]').forEach(r=>r.checked=false); const t=c.querySelector('input.other-note'); t.value=''; t.hidden=true;}); count();});
  document.getElementById('clear-decisions').addEventListener('click',()=>{D={};save(D);cards.forEach(c=>{c.querySelectorAll('input[type=radio]').forEach(r=>r.checked=false);const t=c.querySelector('input.other-note'); t.value=''; t.hidden=true;});count();});
})();
</script>""" % json.dumps(REPO)
    body = howto + "".join(sections) + submit_bar + script
    return {"file": "review.html", "title": "판단 대기 · SKL", "hero": hero("review.html", hero_meta),
            "body": body, "current": "review.html", "asof": ""}


def build(out: Path) -> list[dict]:
    """ai_coop.json·review.json·ai_coop_events.xlsx 를 쓰고 세 페이지(Dashboard·Records·Review)의 사양을 돌려준다(HTML 은 publish_site 가 쓴다)."""
    events, meta = load_events()
    agg = aggregate(events)
    (out / "ai_coop.json").write_text(json.dumps({"as_of": meta["as_of"], "n": agg["n"], "events": events, "agg": agg, "orgs": load_orgs()},
                                                 ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    props = pending_proposals()
    raw_by_id = {e["id"]: e for e in json.loads(EVENTS.read_text(encoding="utf-8")).get("events", [])} if EVENTS.is_file() else {}
    (out / "review.json").write_text(json.dumps({"pending": len(props), "by_task": {k: sum(1 for p in props if p["task"] == k) for k in "ABC"}}, ensure_ascii=False), encoding="utf-8")
    n_xlsx = write_xlsx(out, meta["as_of"])
    print(f"[ai_coop_site] 거래 {agg['n']}건 (확정 {agg['status'].get('confirmed', 0)}, 자동 {agg['status'].get('auto', 0)}) · Excel {n_xlsx}행")
    return [render(agg, {"fetched_at": meta["as_of"]}),
            render_records(agg, {"fetched_at": meta["as_of"]}),
            render_review(props, raw_by_id)]   # 원 레코드(track·sectors) 기준


if __name__ == "__main__":
    print(__doc__)
